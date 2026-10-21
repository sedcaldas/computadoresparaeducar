import os
import sys
import datetime
import docx
import openpyxl
from typing import Dict, Any, Optional
from database import get_db_connection
from drive_service import drive_service

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_026 = os.path.join(BASE_DIR, "LOG-026-Acta de Verificacion y Legalizacion de para bolsa_V1.docx")
TEMPLATE_029 = os.path.join(BASE_DIR, "LOG-029-F Formato Informe Validación Bolsa Tecnologica.xlsx")
TEMPLATE_007 = os.path.join(BASE_DIR, "LOG-007-F LEGALIZACIO_N Y ENTREGA DIRECTA DE EQUIPOS.docx")

TEMP_DIR = os.path.join(BASE_DIR, "generated_documents", "temp")
os.makedirs(TEMP_DIR, exist_ok=True)

class DocumentGenerator:
    """
    Generador oficial de documentos institucionales según Sección 8 de SKILL.md:
    1. LOG-026-F (.docx): Acta de verificación y legalización Bolsa Tecnológica
    2. LOG-029-F (.xlsx): Informe de validación Bolsa Tecnológica
    3. LOG-007-F (.docx): Acta de entrega directa de equipos a estudiantes
    """

    @staticmethod
    def generate_log_026(dane_sede: str, version: int = 1, estado: str = "Borrador") -> Dict[str, Any]:
        """
        Genera LOG-026-F (.docx) para una sede específica con elementos de la Bolsa Tecnológica.
        """
        conn = get_db_connection()
        c = conn.cursor()

        # Obtener datos de sede, institución y municipio
        c.execute("""
            SELECT s.codigo_dane_sede, s.nombre as nombre_sede, s.radicado, s.zona,
                   i.codigo_dane as dane_ie, i.nombre as nombre_ie,
                   m.codigo_municipio, m.nombre as nombre_mpio, m.departamento
            FROM sedes s
            JOIN instituciones i ON s.codigo_dane_ie = i.codigo_dane
            JOIN municipios m ON i.codigo_municipio = m.codigo_municipio
            WHERE s.codigo_dane_sede = ?
        """, (dane_sede,))
        sede_info = c.fetchone()
        if not sede_info:
            conn.close()
            raise ValueError(f"Sede con código DANE {dane_sede} no encontrada.")

        # Obtener elementos de bolsa
        c.execute("""
            SELECT elemento_nombre, cantidad
            FROM bolsa_elementos
            WHERE codigo_dane_sede = ? AND cantidad > 0
            ORDER BY cantidad DESC, elemento_nombre ASC
        """, (dane_sede,))
        elementos = c.fetchall()

        # Cargar plantilla oficial Word
        doc = docx.Document(TEMPLATE_026)
        table = doc.tables[0]

        # R1: Nombre IE & Radicado
        table.rows[1].cells[1].text = str(sede_info["nombre_ie"])
        table.rows[1].cells[5].text = str(sede_info["radicado"] or "")

        # R2: Nombre Sede & DANE Sede
        table.rows[2].cells[1].text = str(sede_info["nombre_sede"])
        table.rows[2].cells[5].text = str(sede_info["codigo_dane_sede"])

        # R3: Departamento & Municipio
        table.rows[3].cells[1].text = str(sede_info["departamento"] or "Caldas")
        table.rows[3].cells[5].text = str(sede_info["nombre_mpio"])

        # Llenar detalle de elementos (R7 a R11)
        # R6 es encabezado: Elemento 1 (col 0), Cant 1 (col 1), Elemento 2 (col 3), Cant 2 (col 5)
        elem_idx = 0
        for r_idx in range(7, 12):
            if elem_idx < len(elementos):
                e1 = elementos[elem_idx]
                table.rows[r_idx].cells[0].text = str(e1["elemento_nombre"])[:60]
                table.rows[r_idx].cells[1].text = str(e1["cantidad"])
                elem_idx += 1
            else:
                table.rows[r_idx].cells[0].text = "N/A"
                table.rows[r_idx].cells[1].text = "0"

            if elem_idx < len(elementos):
                e2 = elementos[elem_idx]
                table.rows[r_idx].cells[3].text = str(e2["elemento_nombre"])[:60]
                table.rows[r_idx].cells[5].text = str(e2["cantidad"])
                elem_idx += 1
            else:
                table.rows[r_idx].cells[3].text = "N/A"
                table.rows[r_idx].cells[5].text = "0"

        # Nombre de archivo
        filename = f"LOG-026-F_{sede_info['codigo_dane_sede']}_v{version:02d}.docx"
        temp_path = os.path.join(TEMP_DIR, filename)
        doc.save(temp_path)

        # Almacenar en la jerarquía de Google Drive
        storage_res = drive_service.store_file(
            src_file_path=temp_path,
            target_filename=filename,
            codigo_mpio=sede_info["codigo_municipio"],
            nombre_mpio=sede_info["nombre_mpio"],
            dane_ie=sede_info["dane_ie"],
            nombre_ie=sede_info["nombre_ie"],
            dane_sede=sede_info["codigo_dane_sede"],
            nombre_sede=sede_info["nombre_sede"],
            tipo_documento="LOG-026-F",
            estado=estado
        )

        # Actualizar o insertar en tabla documentos
        c.execute("""
            SELECT id FROM documentos
            WHERE tipo_documento = 'LOG-026-F' AND codigo_dane_sede = ?
        """, (dane_sede,))
        existing = c.fetchone()
        
        if existing:
            doc_id = existing["id"]
            c.execute("""
                UPDATE documentos
                SET version_actual = ?, estado = ?, ruta_archivo = ?, drive_id = ?,
                    drive_folder_id = ?, hash_sha256 = ?, radicado = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"],
                  storage_res["drive_folder_id"], storage_res["sha256"], sede_info["radicado"], doc_id))
        else:
            c.execute("""
                INSERT INTO documentos (
                    tipo_documento, codigo_dane_sede, radicado, version_actual, estado,
                    ruta_archivo, drive_id, drive_folder_id, hash_sha256
                ) VALUES ('LOG-026-F', ?, ?, ?, ?, ?, ?, ?, ?)
            """, (dane_sede, sede_info["radicado"], version, estado, storage_res["dest_file_path"],
                  storage_res["drive_file_id"], storage_res["drive_folder_id"], storage_res["sha256"]))
            doc_id = c.lastrowid

        # Insertar versión
        c.execute("""
            INSERT INTO documento_versiones (documento_id, version_num, estado, ruta_archivo, drive_id, comentarios)
            VALUES (?, ?, ?, ?, ?, 'Generación automática desde Bolsa Tecnológica')
        """, (doc_id, version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"]))

        conn.commit()
        conn.close()

        return {
            "documento_id": doc_id,
            "filename": filename,
            "path": storage_res["dest_file_path"],
            "drive_id": storage_res["drive_file_id"],
            "drive_path": storage_res["relative_path"],
            "sha256": storage_res["sha256"],
            "version": version
        }

    @staticmethod
    def generate_log_029(dane_sede: str, version: int = 1, estado: str = "Borrador") -> Dict[str, Any]:
        """
        Diligencia LOG-029-F (.xlsx) a partir de la plantilla y datos de la sede,
        preservando fórmulas, hojas auxiliares y catálogos.
        """
        conn = get_db_connection()
        c = conn.cursor()

        c.execute("""
            SELECT s.codigo_dane_sede, s.nombre as nombre_sede, s.radicado,
                   i.codigo_dane as dane_ie, i.nombre as nombre_ie,
                   m.codigo_municipio, m.nombre as nombre_mpio, m.departamento
            FROM sedes s
            JOIN instituciones i ON s.codigo_dane_ie = i.codigo_dane
            JOIN municipios m ON i.codigo_municipio = m.codigo_municipio
            WHERE s.codigo_dane_sede = ?
        """, (dane_sede,))
        sede_info = c.fetchone()
        if not sede_info:
            conn.close()
            raise ValueError(f"Sede con código DANE {dane_sede} no encontrada.")

        c.execute("""
            SELECT elemento_nombre, cantidad
            FROM bolsa_elementos
            WHERE codigo_dane_sede = ? AND cantidad > 0
            ORDER BY cantidad DESC, elemento_nombre ASC
        """, (dane_sede,))
        elementos = c.fetchall()

        wb = openpyxl.load_workbook(TEMPLATE_029, data_only=False)
        ws = wb["Formulario"]

        today_str = datetime.date.today().strftime("%d/%m/%Y")

        # Celdas del formulario oficial
        ws["E8"] = today_str             # Fecha inicial
        ws["I8"] = today_str             # Fecha final
        ws["E9"] = sede_info["departamento"] or "Caldas"
        ws["I9"] = sede_info["nombre_mpio"]
        ws["E10"] = sede_info["nombre_ie"]
        ws["I10"] = sede_info["nombre_sede"]
        ws["E11"] = str(sede_info["radicado"] or "")
        ws["I11"] = str(sede_info["codigo_dane_sede"])

        # Detalle de elementos (a partir de la fila 18)
        row_curr = 18
        for elem in elementos:
            if row_curr > 28:
                break
            ws[f"D{row_curr}"] = str(elem["elemento_nombre"])
            ws[f"E{row_curr}"] = elem["cantidad"]
            ws[f"F{row_curr}"] = "VALIDADO-ZONA-001"
            ws[f"G{row_curr}"] = "SI"
            ws[f"H{row_curr}"] = "SI"
            ws[f"I{row_curr}"] = "Validado sin novedades"
            ws[f"J{row_curr}"] = "Conforme a despacho"
            row_curr += 1

        filename = f"LOG-029-F_{sede_info['codigo_dane_sede']}_v{version:02d}.xlsx"
        temp_path = os.path.join(TEMP_DIR, filename)
        wb.save(temp_path)

        storage_res = drive_service.store_file(
            src_file_path=temp_path,
            target_filename=filename,
            codigo_mpio=sede_info["codigo_municipio"],
            nombre_mpio=sede_info["nombre_mpio"],
            dane_ie=sede_info["dane_ie"],
            nombre_ie=sede_info["nombre_ie"],
            dane_sede=sede_info["codigo_dane_sede"],
            nombre_sede=sede_info["nombre_sede"],
            tipo_documento="LOG-029-F",
            estado=estado
        )

        c.execute("""
            SELECT id FROM documentos
            WHERE tipo_documento = 'LOG-029-F' AND codigo_dane_sede = ?
        """, (dane_sede,))
        existing = c.fetchone()
        
        if existing:
            doc_id = existing["id"]
            c.execute("""
                UPDATE documentos
                SET version_actual = ?, estado = ?, ruta_archivo = ?, drive_id = ?,
                    drive_folder_id = ?, hash_sha256 = ?, radicado = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"],
                  storage_res["drive_folder_id"], storage_res["sha256"], sede_info["radicado"], doc_id))
        else:
            c.execute("""
                INSERT INTO documentos (
                    tipo_documento, codigo_dane_sede, radicado, version_actual, estado,
                    ruta_archivo, drive_id, drive_folder_id, hash_sha256
                ) VALUES ('LOG-029-F', ?, ?, ?, ?, ?, ?, ?, ?)
            """, (dane_sede, sede_info["radicado"], version, estado, storage_res["dest_file_path"],
                  storage_res["drive_file_id"], storage_res["drive_folder_id"], storage_res["sha256"]))
            doc_id = c.lastrowid

        c.execute("""
            INSERT INTO documento_versiones (documento_id, version_num, estado, ruta_archivo, drive_id, comentarios)
            VALUES (?, ?, ?, ?, ?, 'Generación automática informe LOG-029-F')
        """, (doc_id, version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"]))

        conn.commit()
        conn.close()

        return {
            "documento_id": doc_id,
            "filename": filename,
            "path": storage_res["dest_file_path"],
            "drive_id": storage_res["drive_file_id"],
            "drive_path": storage_res["relative_path"],
            "sha256": storage_res["sha256"],
            "version": version
        }

    @staticmethod
    def generate_log_007(documento_estudiante: str, version: int = 1, estado: str = "Generado") -> Dict[str, Any]:
        """
        Genera LOG-007-F (.docx) para un estudiante beneficiario con sus datos,
        acudiente, sede, y equipo asignado, dejando listo el espacio para firma externa.
        """
        conn = get_db_connection()
        c = conn.cursor()

        c.execute("""
            SELECT e.documento_identidad, e.nombres, e.apellidos, e.nombre_completo,
                   e.acudiente_documento, e.acudiente_nombres, e.acudiente_parentesco, e.acudiente_telefono,
                   s.codigo_dane_sede, s.nombre as nombre_sede, s.radicado,
                   i.codigo_dane as dane_ie, i.nombre as nombre_ie,
                   m.codigo_municipio, m.nombre as nombre_mpio, m.departamento,
                   a.tipo_terminal, a.serial_asignado
            FROM estudiantes e
            JOIN sedes s ON e.codigo_dane_sede = s.codigo_dane_sede
            JOIN instituciones i ON s.codigo_dane_ie = i.codigo_dane
            JOIN municipios m ON i.codigo_municipio = m.codigo_municipio
            LEFT JOIN asignaciones a ON e.documento_identidad = a.estudiante_documento
            WHERE e.documento_identidad = ?
        """, (documento_estudiante,))
        est_info = c.fetchone()
        if not est_info:
            conn.close()
            raise ValueError(f"Estudiante con documento {documento_estudiante} no encontrado.")

        doc = docx.Document(TEMPLATE_007)
        table_basic = doc.tables[0]
        table_firms = doc.tables[2]

        today_str = datetime.date.today().strftime("%d de %B de %Y")

        # Table 0:
        # R1: Fecha de legalización
        table_basic.rows[1].cells[1].text = today_str
        # R2: Nombre de IE y Sede | Código DANE Sede
        table_basic.rows[2].cells[1].text = f"{est_info['nombre_ie']} - {est_info['nombre_sede']}"
        table_basic.rows[2].cells[3].text = str(est_info["codigo_dane_sede"])
        # R3: Nombre del rector
        table_basic.rows[3].cells[1].text = f"Rector(a) {est_info['nombre_ie']}"
        # R4: Cantidad de folios
        table_basic.rows[4].cells[1].text = "1 (Un folio)"
        # R5: Radicado Sede
        table_basic.rows[5].cells[1].text = str(est_info["radicado"] or "")
        # R6: Departamento y Municipio
        table_basic.rows[6].cells[1].text = str(est_info["departamento"] or "Caldas")
        table_basic.rows[6].cells[3].text = str(est_info["nombre_mpio"])
        # R7: Tipo de terminal
        tipo_term = est_info["tipo_terminal"] or "Tableta"
        if tipo_term.lower() == "tableta":
            table_basic.rows[7].cells[1].text = "● Tabletas       ○ Portátiles    ○ Otro"
        else:
            table_basic.rows[7].cells[1].text = "○ Tabletas       ● Portátiles    ○ Otro"
        # R8: Cantidad de terminales
        table_basic.rows[8].cells[1].text = "1 (Uno)"

        # Table 2: Lista de estudiantes / acudientes
        # R0 es encabezado. R1 es la primera fila de datos.
        table_firms.rows[1].cells[0].text = str(est_info["documento_identidad"])
        table_firms.rows[1].cells[1].text = str(est_info["nombre_completo"])
        table_firms.rows[1].cells[2].text = str(est_info["acudiente_nombres"] or "")
        table_firms.rows[1].cells[3].text = str(est_info["acudiente_documento"] or "")
        table_firms.rows[1].cells[4].text = str(est_info["acudiente_telefono"] or "3100000000")
        table_firms.rows[1].cells[5].text = ""  # Espacio para firma física

        filename = f"LOG-007-F_{est_info['codigo_dane_sede']}_{est_info['documento_identidad']}_v{version:02d}.docx"
        temp_path = os.path.join(TEMP_DIR, filename)
        doc.save(temp_path)

        storage_res = drive_service.store_file(
            src_file_path=temp_path,
            target_filename=filename,
            codigo_mpio=est_info["codigo_municipio"],
            nombre_mpio=est_info["nombre_mpio"],
            dane_ie=est_info["dane_ie"],
            nombre_ie=est_info["nombre_ie"],
            dane_sede=est_info["codigo_dane_sede"],
            nombre_sede=est_info["nombre_sede"],
            tipo_documento="LOG-007-F",
            estado=estado,
            documento_estudiante=est_info["documento_identidad"]
        )

        c.execute("""
            SELECT id FROM documentos
            WHERE tipo_documento = 'LOG-007-F' AND estudiante_documento = ?
        """, (documento_estudiante,))
        existing = c.fetchone()
        
        if existing:
            doc_id = existing["id"]
            c.execute("""
                UPDATE documentos
                SET version_actual = ?, estado = ?, ruta_archivo = ?, drive_id = ?,
                    drive_folder_id = ?, hash_sha256 = ?, radicado = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
            """, (version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"],
                  storage_res["drive_folder_id"], storage_res["sha256"], est_info["radicado"], doc_id))
        else:
            c.execute("""
                INSERT INTO documentos (
                    tipo_documento, codigo_dane_sede, estudiante_documento, radicado,
                    version_actual, estado, ruta_archivo, drive_id, drive_folder_id, hash_sha256
                ) VALUES ('LOG-007-F', ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (est_info["codigo_dane_sede"], documento_estudiante, est_info["radicado"],
                  version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"],
                  storage_res["drive_folder_id"], storage_res["sha256"]))
            doc_id = c.lastrowid

        c.execute("""
            INSERT INTO documento_versiones (documento_id, version_num, estado, ruta_archivo, drive_id, comentarios)
            VALUES (?, ?, ?, ?, ?, 'Generación inicial formato LOG-007-F para firma')
        """, (doc_id, version, estado, storage_res["dest_file_path"], storage_res["drive_file_id"]))

        conn.commit()
        conn.close()

        return {
            "documento_id": doc_id,
            "filename": filename,
            "path": storage_res["dest_file_path"],
            "drive_id": storage_res["drive_file_id"],
            "drive_path": storage_res["relative_path"],
            "sha256": storage_res["sha256"],
            "version": version
        }
